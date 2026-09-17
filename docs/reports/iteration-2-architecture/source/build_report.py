"""Build the editable report and its Markdown copy using the bundled Python runtime."""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from report_content import PAGES

ROOT = Path(__file__).resolve().parents[1]
STEM = 'AI-Wrevolusi-Iteration-2-System-Architecture-Report'
NAVY = '17324D'
TEAL = '166A72'
GRAY = '536575'


def shade(cell, fill):
    item = OxmlElement('w:shd')
    item.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(item)


def format_section(section, landscape):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11.6929 if landscape else 8.2677)
    section.page_height = Inches(8.2677 if landscape else 11.6929)
    section.top_margin = Inches(0.5 if landscape else 0.66)
    section.bottom_margin = Inches(0.5 if landscape else 0.62)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.23)


def add_table(doc, block):
    widths = block.get('widths') or [6.7 / len(block['headers'])] * len(block['headers'])
    result = doc.add_table(rows=1, cols=len(block['headers']))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for index, title in enumerate(block['headers']):
        result.rows[0].cells[index].text = title
    repeat = OxmlElement('w:tblHeader')
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for data in block['rows']:
        for cell, text in zip(result.add_row().cells, data):
            cell.text = text
    for r, row in enumerate(result.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for c, cell in enumerate(row.cells):
            cell.width = Inches(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '90')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '4')
                e.set(qn('w:color'), 'D7E0E7')
                borders.append(e)
            props.append(borders)
            shade(cell, NAVY if r == 0 else ('F3F6F8' if r % 2 else 'FFFFFF'))
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.font.bold = r == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if r == 0 else NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    document = Document()
    styles = document.styles
    for name in ['Normal', 'Body Text', 'List Bullet']:
        styles[name].font.name = 'Arial'
        styles[name].font.size = Pt(10.5)
        styles[name].paragraph_format.space_after = Pt(7)
        styles[name].paragraph_format.line_spacing = 1.1
    for name, size, color in [('Title', 28, NAVY), ('Heading 1', 20, NAVY), ('Heading 2', 12, TEAL), ('Caption', 9.5, GRAY)]:
        styles[name].font.name = 'Arial'
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)
        styles[name].font.bold = name != 'Caption'
        styles[name].paragraph_format.space_before = Pt(9 if name == 'Heading 2' else 0)
        styles[name].paragraph_format.space_after = Pt(7)
        styles[name].paragraph_format.keep_with_next = name != 'Caption'
    styles['Caption'].font.italic = False
    props = document.core_properties
    props.title = 'AI-Wrevolusi - Iteration 2 System Architecture Report'
    props.subject = 'Architecture evolution, design reasons and current data model'
    props.author = 'AI-Wrevolusi project team'
    props.keywords = 'AI-Wrevolusi, Iteration 2, system architecture, ERD'
    props.created = props.modified = datetime(2026, 9, 17, 14, 27, 41, tzinfo=timezone.utc)
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('AI-Wrevolusi | Iteration 2 | ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    markdown = []
    for index, page in enumerate(PAGES):
        landscape = page['landscape']
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        format_section(section, landscape)
        title = document.add_paragraph(page['title'], style='Title' if index == 0 else 'Heading 1')
        if landscape:
            title.runs[0].font.size = Pt(17)
        markdown += ['# ' + page['title'].replace('\n', ' - '), '']
        for block in page['blocks']:
            kind = block['kind']
            if kind in ('p', 'h'):
                assert ';' not in block['text'], block['text']
                para = document.add_paragraph(block['text'], style='Heading 2' if kind == 'h' else 'Normal')
                if index == 0 and block['text'].startswith(('FIT5120', 'Current code')):
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor.from_string(GRAY)
                markdown += [('## ' if kind == 'h' else '') + block['text'], '']
            elif kind == 'bullets':
                for text in block['items']:
                    assert ';' not in text
                    document.add_paragraph(text, style='List Bullet')
                    markdown += ['- ' + text]
                markdown += ['']
            elif kind == 'table':
                assert all(';' not in x for row in [block['headers'], *block['rows']] for x in row)
                add_table(document, block)
                markdown += ['| ' + ' | '.join(block['headers']) + ' |', '| ' + ' | '.join(['---'] * len(block['headers'])) + ' |']
                markdown += ['| ' + ' | '.join(cell.replace('\n', '<br>') for cell in row) + ' |' for row in block['rows']]
                markdown += ['']
            elif kind == 'figure':
                path = ROOT / 'figures' / block['name']
                with Image.open(path) as img:
                    width, height = img.size
                display_width = min(10.25, 5.86 * width / height)
                para = document.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.keep_with_next = True
                image = para.add_run().add_picture(str(path), width=Inches(display_width))
                image._inline.docPr.set('descr', block['caption'])
                document.add_paragraph(block['caption'], style='Caption')
                markdown += [f"![{block['caption']}](figures/{block['name']})", '', block['caption'], '']
    # Remove default Word title rules so the report uses plain heading typography.
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    docx_path = ROOT / (STEM + '.docx')
    document.save(docx_path)
    (ROOT / (STEM + '.md')).write_text('\n'.join(markdown))
    print(docx_path)
    print('Planned pages:', len(PAGES))

if __name__ == '__main__':
    build()
